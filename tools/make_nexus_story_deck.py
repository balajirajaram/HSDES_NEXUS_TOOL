"""Generate a fresh NEXUS executive story deck.

Presentation-only artifact generator. It does not alter application or
qualification logic.
"""
from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation"
PPTX = OUT / "NEXUS_Executive_Story.pptx"
ARCH = OUT / "NEXUS_Executive_Architecture.png"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
BLUE = RGBColor(0x00, 0x71, 0xC5)
TEAL = RGBColor(0x00, 0x85, 0x7D)
AMBER = RGBColor(0xB2, 0x6A, 0x00)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xB3, 0x26, 0x1E)
INK = RGBColor(0x26, 0x36, 0x4F)
GREY = RGBColor(0x5B, 0x6B, 0x7F)
PALE_BLUE = RGBColor(0xE9, 0xF4, 0xFB)
PALE_TEAL = RGBColor(0xE6, 0xF3, 0xF1)
PALE_GREEN = RGBColor(0xE9, 0xF5, 0xEA)
PALE_AMBER = RGBColor(0xFF, 0xF3, 0xDC)
PALE_RED = RGBColor(0xFC, 0xEB, 0xE9)
PALE_GREY = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return box


def header(slide, title, subtitle, number, prs):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    add_text(slide, title, 0.55, 0.17, 12.1, 0.42, 27, WHITE, True)
    add_text(slide, subtitle, 0.58, 0.68, 12.0, 0.22, 11, RGBColor(0xC9, 0xD6, 0xEA))
    add_text(slide, f"Intel Confidential  |  {number}", 0.55, 7.12, 12, 0.18, 9, GREY)


def card(slide, x, y, w, h, title, body, fill=PALE_BLUE, edge=BLUE, title_size=15, body_size=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = edge; shape.line.width = Pt(1)
    add_text(slide, title, x + 0.18, y + 0.16, w - 0.36, 0.3, title_size, edge, True)
    add_text(slide, body, x + 0.18, y + 0.58, w - 0.36, h - 0.7, body_size, INK)


def tile(slide, x, y, w, label, value, detail, fill, edge):
    card(slide, x, y, w, 1.45, label, detail, fill, edge, 10, 10)
    add_text(slide, value, x + 0.18, y + 0.46, w - 0.36, 0.42, 24, edge, True)


def bullets(slide, items, x=0.8, y=1.55, w=11.8, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.3))
    tf = box.text_frame; tf.word_wrap = True
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text; p.font.size = Pt(size); p.font.color.rgb = INK; p.space_after = Pt(12)
    return box


def flow(slide, labels, y=3.0):
    gap = 0.14; x = 0.45; w = (12.4 - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.92))
        shape.fill.solid(); shape.fill.fore_color.rgb = PALE_BLUE
        shape.line.color.rgb = BLUE
        add_text(slide, label, x + 0.07, y + 0.22, w - 0.14, 0.46, 11, NAVY, True, PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            add_text(slide, ">", x + w + 0.02, y + 0.28, gap - 0.04, 0.3, 18, BLUE, True, PP_ALIGN.CENTER)
        x += w + gap


def build_architecture_png():
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    fig, ax = plt.subplots(figsize=(16, 8.5), dpi=150)
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.add_patch(FancyBboxPatch((0, .9), 1, .1, boxstyle="square", facecolor="#1B2A4A", edgecolor="none"))
    ax.text(.5, .95, "NEXUS — Evidence-graded HSDES triage", ha="center", va="center", color="white", fontsize=22, weight="bold")
    boxes = [
        (.04, .56, .16, .18, "HSDES", "Ticket\ncomments\nattachments", "#E9F4FB", "#0071C5"),
        (.25, .56, .16, .18, "Evidence", "MCA\nfirst error\nboot flow", "#E6F3F1", "#00857D"),
        (.46, .56, .16, .18, "Ownership", "reporting\noriginating\nvictim IP", "#FFF3DC", "#B26A00"),
        (.67, .56, .16, .18, "Judgment", "confidence\nverdict\nmissing data", "#FFF3DC", "#B26A00"),
        (.85, .56, .11, .18, "Gate", "draft\nor post", "#E9F5EA", "#2E7D32"),
    ]
    for x, y, w, h, title, body, fill, edge in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=.01", facecolor=fill, edgecolor=edge, linewidth=2))
        ax.text(x+w/2, y+h-.04, title, ha="center", va="top", color="#1B2A4A", fontsize=13, weight="bold")
        ax.text(x+w/2, y+h-.1, body, ha="center", va="top", color="#26364F", fontsize=10, linespacing=1.5)
    for x in (.2, .41, .62, .83):
        ax.add_patch(FancyArrowPatch((x, .65), (x+.045, .65), arrowstyle="-|>", mutation_scale=15, color="#0071C5", linewidth=2))
    ax.add_patch(FancyBboxPatch((.22, .2), .56, .15, boxstyle="round,pad=.01", facecolor="#F3F5F7", edgecolor="#5B6B7F", linewidth=1.5))
    ax.text(.5, .29, "Safety boundary", ha="center", va="center", fontsize=15, color="#1B2A4A", weight="bold")
    ax.text(.5, .235, "Human claims are context — not proof.  Contradictions and ambiguity block auto-post.", ha="center", va="center", fontsize=11, color="#26364F")
    fig.savefig(ARCH, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_deck():
    OUT.mkdir(exist_ok=True)
    build_architecture_png()
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    n = 0

    def make(title, subtitle):
        nonlocal n
        n += 1
        slide = prs.slides.add_slide(blank)
        header(slide, title, subtitle, n, prs)
        return slide

    slide = prs.slides.add_slide(blank); n += 1
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    add_text(slide, "NEXUS", 0.85, 1.8, 5.0, .8, 48, WHITE, True)
    add_text(slide, "Evidence-graded HSDES triage for Xeon validation", 0.9, 2.8, 10.8, .55, 25, RGBColor(0xC9, 0xD6, 0xEA))
    add_text(slide, "A practical bridge from ticket noise to defensible engineering action", 0.9, 3.75, 10.4, .45, 19, WHITE)
    add_text(slide, "Intel Confidential  |  2026-09-11", .9, 6.65, 10, .3, 12, RGBColor(0xC9, 0xD6, 0xEA))

    slide = make("The Debugging Problem", "Why a ticket ID is not yet an explanation")
    tile(slide, .7, 1.65, 2.8, "INPUT", "HSDES", "Narrative, comments, attachments, partial logs", PALE_BLUE, BLUE)
    tile(slide, 3.75, 1.65, 2.8, "RISK", "Noise", "Many signals are symptoms, victims, or snapshots", PALE_RED, RED)
    tile(slide, 6.8, 1.65, 2.8, "COST", "Rework", "Known signatures are re-investigated across teams", PALE_AMBER, AMBER)
    tile(slide, 9.85, 1.65, 2.8, "NEED", "Evidence", "A traceable route from machine facts to action", PALE_TEAL, TEAL)
    bullets(slide, ["The hard problem is not producing prose; it is deciding what the evidence actually proves.", "A human comment can be valuable context without being machine proof.", "A reporting bank, first-error source, and originating owner are not automatically the same thing."], .9, 4.05, 11.4, 19)

    slide = make("The NEXUS Answer", "A disciplined evidence path")
    flow(slide, ["Read", "Decode", "Separate", "Grade", "Gate", "Act"], 2.75)
    bullets(slide, ["Read the ticket and attached evidence.", "Decode MCA, boot, serial, PythonSV, and first-error signals.", "Separate machine evidence from human claims and historical recall.", "Grade ownership, confidence, contradictions, and missing evidence.", "Block weak or ambiguous conclusions from autonomous posting.", "Return a report, draft, or explicitly bounded next action."], .9, 4.45, 11.4, 15)

    slide = make("Architecture", "Where evidence changes shape")
    slide.shapes.add_picture(str(ARCH), Inches(.45), Inches(1.35), width=Inches(12.45))

    slide = make("Safety Is the Product", "The system is designed to know when not to overclaim")
    tile(slide, .7, 1.65, 3.7, "COMMENT ISOLATION", "No", "Comment-only root cause cannot become CONFIRMED", PALE_GREEN if False else PALE_TEAL, TEAL)
    tile(slide, 4.8, 1.65, 3.7, "MCA FATALITY", "Never", "UC + PCC is never mislabeled corrected", PALE_TEAL, TEAL)
    tile(slide, 8.9, 1.65, 3.7, "POSTING", "Blocked", "Conflict, ambiguity, and weak evidence stay drafts", PALE_RED, RED)
    bullets(slide, ["Safety invariants are explicit, named, and regression-tested.", "HSDES_WRITE_ENABLED defaults false.", "force=true cannot bypass write-disable."], .95, 4.55, 11.3, 19)

    slide = make("AutoHSD: From Ticket to Evidence", "Closed-loop collection with a narrow safety boundary")
    flow(slide, ["HSD", "Evidence\naudit", "Reachability", "Read-only\nSSH/BMC", "Re-analyze", "Gate", "Draft/Post"], 2.9)
    add_text(slide, "Collection is conditional. Existing attachments win; unreachable nodes produce insufficient evidence rather than fabricated logs.", .95, 5.1, 11.4, .65, 18, NAVY, True)

    slide = make("What Ships Today", "A focused engineering tool, not a promise of autonomous closure")
    tile(slide, .7, 1.7, 2.8, "RCA", "Structured", "Evidence, ownership, verdict, missing data", PALE_BLUE, BLUE)
    tile(slide, 3.75, 1.7, 2.8, "AUTOHSD", "Guarded", "SSH/BMC collection and before/after analysis", PALE_TEAL, TEAL)
    tile(slide, 6.8, 1.7, 2.8, "KNOWLEDGE", "Traceable", "KB, provenance, historical reference", PALE_BLUE, BLUE)
    tile(slide, 9.85, 1.7, 2.8, "WRITE-BACK", "Draft-first", "Post only after configuration and gate", PALE_AMBER, AMBER)
    bullets(slide, ["The tool can say: insufficient evidence, contradiction, or working hypothesis.", "That behavior is a feature for validation teams: uncertainty is visible and actionable."], .95, 4.55, 11.3, 20)

    slide = make("Current Evidence Snapshot", "Numbers from the current workspace, not a marketing baseline")
    tile(slide, .7, 1.7, 2.8, "REGRESSION", "82/82", "Current full test suite", PALE_GREEN, GREEN)
    tile(slide, 3.75, 1.7, 2.8, "PROVENANCE", "18/18", "Inventory coverage; 0 unknown", PALE_GREEN, GREEN)
    tile(slide, 6.8, 1.7, 2.8, "GOLDEN", "47", "Strict Level 3/4 qualified cases", PALE_BLUE, BLUE)
    tile(slide, 9.85, 1.7, 2.8, "AUTOHSD", "10/10", "Focused integration checks", PALE_GREEN, GREEN)
    add_text(slide, "Production qualification remains open: benchmark sanity has actionable predictions, but this is not yet a completed accuracy qualification.", .95, 4.75, 11.3, .8, 19, AMBER, True)

    slide = make("A Concrete Safety Case", "16031734105: reporting owner versus first-error source")
    card(slide, .8, 1.7, 3.65, 2.0, "Naive path", "Select the reporting CCF bank and state a confident CCF root cause.", PALE_RED, RED, 16, 16)
    card(slide, 4.85, 1.7, 3.65, 2.0, "NEXUS path", "Preserve PUNIT first-error evidence, CCF reporting evidence, and decoder ambiguity.", PALE_TEAL, TEAL, 16, 16)
    card(slide, 8.9, 1.7, 3.65, 2.0, "Outcome", "WORKING HYPOTHESIS; contradiction detector; auto-post blocked.", PALE_BLUE, BLUE, 16, 16)
    add_text(slide, "The value is not a louder answer. It is a more defensible boundary around the answer.", 1.0, 5.0, 11.1, .5, 20, NAVY, True)

    slide = make("Knowledge That Stays Honest", "Mining accelerates review; validation creates trust")
    tile(slide, .8, 1.7, 3.5, "HISTORICAL GRAPH", "88", "Nodes", PALE_BLUE, BLUE)
    tile(slide, 4.9, 1.7, 3.5, "RELATIONSHIPS", "421", "Edges", PALE_BLUE, BLUE)
    tile(slide, 9.0, 1.7, 3.5, "CANDIDATES", "88", "Consensus rows for review", PALE_AMBER, AMBER)
    bullets(slide, ["Candidate mining is not Golden truth.", "Bugeco, sighting confirmation, reproduction, and engineer review determine validation tier.", "The corpus remains a controlled evidence set, not a random similarity dump."], .95, 4.6, 11.3, 18)

    slide = make("What the Engineer Sees", "A report that is useful before it is certain")
    card(slide, .75, 1.65, 3.75, 2.05, "Executive summary", "What happened, what is known, what is missing.", PALE_BLUE, BLUE, 16, 16)
    card(slide, 4.8, 1.65, 3.75, 2.05, "Evidence chain", "MCA, first error, boot stage, ownership, contradictions.", PALE_TEAL, TEAL, 16, 16)
    card(slide, 8.85, 1.65, 3.75, 2.05, "Next action", "The smallest useful collection or validation step.", PALE_AMBER, AMBER, 16, 16)
    add_text(slide, "Historical reproduction suggestions are explicitly reference-only and do not change confidence, verdict, or posting.", .95, 5.1, 11.3, .5, 18, GREY, True)

    slide = make("What Is Not Claimed", "Boundaries matter")
    bullets(slide, ["Not a replacement for silicon, firmware, or validation engineering judgment.", "Not a guarantee that every ticket has enough evidence for an RCA.", "Not an autonomous writer by default.", "Not a substitute for the centralized Elastic log pipeline; direct collection is limited to configured SSH/BMC paths.", "Not production-qualified until the evidence-backed benchmark is completed."], .95, 1.7, 11.3, 20)

    slide = make("Next Moves", "The shortest path to a defensible pilot")
    bullets(slide, ["Complete evidence-backed benchmark validation with representative workloads and raw logs.", "Shadow-mode AutoHSD on a small volunteer cohort.", "Measure time-to-first-useful-action, not just report generation time.", "Expand domain review for borderline ownership and decoder cases.", "Only then decide whether shared deployment and broader operations are warranted."], .95, 1.7, 11.3, 20)

    slide = make("The Ask", "A focused collaboration request")
    card(slide, 1.0, 1.8, 3.5, 2.2, "Validation leaders", "Nominate representative HSDs and review borderline cases.", PALE_BLUE, BLUE, 17, 17)
    card(slide, 4.9, 1.8, 3.5, 2.2, "Debug experts", "Confirm evidence ownership and resolution tiers.", PALE_TEAL, TEAL, 17, 17)
    card(slide, 8.8, 1.8, 3.5, 2.2, "Pilot sponsors", "Support a shadow-mode, draft-only evaluation.", PALE_AMBER, AMBER, 17, 17)
    add_text(slide, "The next decision should be based on observed engineering value and evidence quality.", 1.1, 5.2, 11, .45, 20, NAVY, True)

    slide = make("Appendix", "NEXUS in one sentence")
    add_text(slide, "NEXUS connects HSD symptoms to evidence, evidence to ownership, and ownership to a guarded engineering action — while making uncertainty impossible to hide.", 1.1, 2.2, 11.1, 1.4, 27, NAVY, True)
    add_text(slide, "Presentation generated locally from the current repository state.", 1.1, 5.3, 10.5, .35, 14, GREY)

    prs.save(PPTX)
    print(f"PPTX: {PPTX}")
    print(f"PNG: {ARCH}")


if __name__ == "__main__":
    build_deck()
